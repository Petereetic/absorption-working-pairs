#!/usr/bin/env python3
"""Replace the plain-text formulas in the R22-DMF docx with high-definition formula images rendered by matplotlib.
Usage: python3 beautify_formulas_img.py <docx file>
Modifies the file in place; first backed up as .bak.
"""
import sys
import os
import shutil
import re
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Pt, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH

IMG_DIR = '/tmp/formula_imgs'
os.makedirs(IMG_DIR, exist_ok=True)

# (key, mathtext, mode)  mode: display = standalone centred line, inline = inline
FORMULAS = [
    ('antoine', r'$\ln(P/\mathrm{bar}) = A + \dfrac{B}{T} + C\cdot\ln T$', 'inline'),
    ('hs', r'$h_s(T,\xi) = \xi\cdot h_{L,R22}(T) + (1-\xi)\cdot h_{DMF}(T) + h_{\mathrm{mix}}(T,\xi)$', 'display'),
    ('hmix', r'$h_{\mathrm{mix}} = \left[\dfrac{(1-\xi)\cdot R\cdot T^2}{M_{\mathrm{mix}}}\right]\cdot(K_0\cdot Y_0 + K_1\cdot Y_1 + K_2\cdot Y_2 + K_3\cdot Y_3)$', 'display-after'),
    ('y0', r'$Y_0 = \dfrac{\xi}{1-\xi}$', 'display'),
    ('y1', r'$Y_1 = \dfrac{\xi}{1-\xi} + \ln(1-\xi)$', 'display'),
    ('y2', r'$Y_2 = \dfrac{1}{1-\xi} - (1-\xi) + 2\cdot\ln(1-\xi)$', 'display'),
    ('y3', r'$Y_3 = \dfrac{\xi}{1-\xi} + \dfrac{\xi^2}{2} + 2\xi + 3\cdot\ln(1-\xi)$', 'display'),
    ('k0', r'$K_0 = \dfrac{B_0}{T^2} + \dfrac{2C_0}{T^3} - \dfrac{E_1}{T^2} + E_2 + \dfrac{E_3}{T}$', 'display'),
    ('k1', r'$K_1 = \dfrac{B_1}{T^2} + \dfrac{2C_1}{T^3}$', 'display'),
    ('k2', r'$K_2 = \dfrac{B_2}{T^2} + \dfrac{2C_2}{T^3}$', 'display'),
    ('k3', r'$K_3 = 0$', 'inline'),
    ('mmix', r'$M_{\mathrm{mix}}=\xi\cdot M_{R22}+(1-\xi)\cdot M_{DMF}$', 'inline'),
    ('hf', r'$h=F_0+F_1\cdot T+F_2\cdot T^2$', 'inline'),
]

def render_formula(key, mathtext, fontsize=13):
    """Render a single formula to PNG; returns the path."""
    path = os.path.join(IMG_DIR, f'{key}.png')
    fig = plt.figure(figsize=(0.1, 0.1), dpi=300)
    fig.patch.set_alpha(0.0)
    txt = fig.text(0.5, 0.5, mathtext, fontsize=fontsize,
                   ha='center', va='center')
    # Draw once first to compute the bbox
    fig.canvas.draw()
    bbox = txt.get_window_extent(renderer=fig.canvas.get_renderer())
    # Add padding (pixels)
    pad = 12
    w_in = (bbox.width + 2*pad) / 300
    h_in = (bbox.height + 2*pad) / 300
    fig.set_size_inches(w_in, h_in)
    # Re-position to the centre
    txt.set_position((0.5, 0.5))
    fig.savefig(path, dpi=300, transparent=True,
                bbox_inches='tight', pad_inches=0.05)
    plt.close(fig)
    return path

# Paragraph matching rules: (regex, formula key)
RULES = [
    (re.compile(r'ln\(P/bar\) = A \+ B/T \+ C·lnT'), 'antoine'),
    (re.compile(r'^h_s\(T,ξ\) = ξ·h_L,R22\(T\) \+ \(1−ξ\)·h_DMF\(T\) \+ h_mix\(T,ξ\)$'), 'hs'),
    (re.compile(r'h_mix = \[\(1−ξ\)·R·T²/M_mix\]·\(K0·Y0 \+ K1·Y1 \+ K2·Y2 \+ K3·Y3\)'), 'hmix'),
    (re.compile(r'^Y0 = ξ/\(1−ξ\)$'), 'y0'),
    (re.compile(r'^Y1 = ξ/\(1−ξ\) \+ ln\(1−ξ\)$'), 'y1'),
    (re.compile(r'^Y2 = 1/\(1−ξ\) − \(1−ξ\) \+ 2·ln\(1−ξ\)$'), 'y2'),
    (re.compile(r'^Y3 = ξ/\(1−ξ\) \+ ξ²/2 \+ 2ξ \+ 3·ln\(1−ξ\)$'), 'y3'),
    (re.compile(r'^K0 = B0/T² \+ 2C0/T³ − E1/T² \+ E2 \+ E3/T$'), 'k0'),
    (re.compile(r'^K1 = B1/T² \+ 2C1/T³$'), 'k1'),
    (re.compile(r'^K2 = B2/T² \+ 2C2/T³$'), 'k2'),
    (re.compile(r'^K3 = 0'), 'k3'),
    (re.compile(r'M_mix=ξ·M_R22\+\(1−ξ\)·M_DMF'), 'mmix'),
    (re.compile(r'h=F0\+F1·T\+F2·T²'), 'hf'),
]

def clear_paragraph(p):
    pPr = p._p.get_or_add_pPr()
    for child in list(p._p):
        if child is not pPr:
            p._p.remove(child)

def process_docx(path, img_paths, modes):
    shutil.copy(path, path + '.bak')
    doc = Document(path)
    changed = 0
    for p in doc.paragraphs:
        t = p.text
        if not t.strip():
            continue
        for pat, key in RULES:
            m = pat.search(t)
            if not m:
                continue
            mode = modes[key]
            before = t[:m.start()]
            after = t[m.end():]
            clear_paragraph(p)
            if mode == 'display':
                # Standalone line: centred image
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(img_paths[key])
            elif mode == 'display-after':
                # Keep the lead-in word (the "where" connector) in this paragraph; the formula goes on its own centred line
                if before.strip():
                    p.add_run(before.strip())
                # Insert a formula paragraph after the current paragraph
                from docx.oxml import OxmlElement
                new_p = OxmlElement('w:p')
                p._p.addnext(new_p)
                from docx.text.paragraph import Paragraph
                fp = Paragraph(new_p, p._parent)
                fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                fp.add_run().add_picture(img_paths[key])
                if after.strip():
                    # Rare case: text still follows the formula; start a new paragraph
                    new_p2 = OxmlElement('w:p')
                    new_p.addnext(new_p2)
                    Paragraph(new_p2, p._parent).add_run(after.strip())
            else:
                # Inline: text + image + text
                if before:
                    p.add_run(before)
                run = p.add_run()
                # Inline formula height is about 1.2x the line height; use a fixed height
                run.add_picture(img_paths[key], height=Pt(16))
                if after:
                    p.add_run(after)
            changed += 1
            print(f'  replaced: {t[:60]}...')
            break
    doc.save(path)
    print(f'{path}: {changed} replacements in total')

if __name__ == '__main__':
    print('Rendering formula images...')
    img_paths = {}
    modes = {}
    for key, mathtext, mode in FORMULAS:
        img_paths[key] = render_formula(key, mathtext)
        modes[key] = mode
        print(f'  {key}: {img_paths[key]}')
    for f in sys.argv[1:]:
        process_docx(f, img_paths, modes)
