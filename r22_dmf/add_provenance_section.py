# -*- coding: utf-8 -*-
"""在两份 R22-DMF .docx 末尾（适用范围之后、注脚之前）插入"文章是怎么来的"章节。"""
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ZH_PARAS = [
    ("五、这篇文章是怎么来的", True),
    ("说点题外话。渔夫是个业余研究者，没有单位图书馆，做吸收式制冷纯属个人兴趣。"
     "Fatouh 等 1993 年那篇焓关联式的原文，出版商要价几十美元一篇——业余爱好者掏这钱，肉疼。", False),
    ("办法是 Google 搜索。Ardita 2008 年的硕士论文在网上能免费下载，一分钱没花。"
     "可论文不是英文（印尼语，\"Lampiran\" 就是\"附录\"的意思），下载了就放在硬盘里吃灰——"
     "公式看得懂，文字看不懂，数据表也不敢直接拿来用。", False),
    ("转机是 Muse。论文从硬盘里翻出来之后，后台的活儿都是它干的：附录 7 的 132 个数据点逐点转录成表，"
     "253.15 K 那处 0.610 bar 对着扫描原图逐字核对过；10 条等浓度线各拟一条 Antoine 式，"
     "浓度方向用 PCHIP 保单调插值，搭出泡点压力模型；Fatouh 关联式逐行翻译成 Python、C、VBA 三套实现，"
     "C 与 Python 逐点比对最大偏差 5×10⁻⁷；还顺手揪出论文手算例（98.7 °C、ξ=0.5，383.75 kJ/kg）"
     "疑似把摄氏度数值误代入了开尔文公式，按 K 算应为 283.6；四张物性图和这篇中英文文章，也都是它排的版。", False),
    ("渔夫负责拍板和校验，Muse 负责搬砖和算数。原材料免费，AI 工也免费——这篇文章就是这么来的。", False),
]

EN_PARAS = [
    ("5. How this article came about", True),
    ("A word about provenance. The fisherman is an amateur researcher with no institutional library; "
     "absorption refrigeration is a personal interest. The original Fatouh et al. 1993 paper with the "
     "enthalpy correlation costs several tens of dollars from the publisher — real money for a hobbyist.", False),
    ("The workaround was a Google search. Ardita's 2008 master's thesis is freely downloadable — "
     "not a cent spent. But it isn't in English (it's Indonesian; \"Lampiran\" means \"appendix\"), "
     "so it sat on the hard drive gathering dust: the equations were readable, the prose was not, "
     "and the data tables couldn't be used blindly.", False),
    ("Enter Muse. Once the thesis was dug out, it did the background work: transcribing the 132 data "
     "points from Appendix 7, verifying the 0.610 bar at 253.15 K character by character against the "
     "scanned original; fitting one Antoine equation per isoconcentration line and interpolating in "
     "concentration with monotone PCHIP to build the bubble-pressure model; translating the Fatouh "
     "correlation line by line into Python, C, and VBA, with C vs. Python agreeing to 5×10⁻⁷; "
     "catching that the paper's worked example (98.7 °C, ξ=0.5, 383.75 kJ/kg) looks like a Celsius value "
     "mistakenly plugged into a Kelvin formula (283.6 kJ/kg with K); and laying out the four property "
     "charts and these Chinese and English articles.", False),
    ("The fisherman made the calls and checked the results; Muse did the lifting and the arithmetic. "
     "Free raw material, free AI labor — that is how this article came to be.", False),
]


def insert_before(ref_para, text, bold, font_name=None):
    new_el = OxmlElement("w:p")
    ref_para._p.addprevious(new_el)
    p = Paragraph(new_el, ref_para._parent)
    run = p.add_run(text)
    run.bold = bold
    if font_name:
        run.font.name = font_name
        rPr = run._r.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        rFonts.set(qn("w:eastAsia"), font_name)
    if bold:
        run.font.size = Pt(14)
    return p


def add_section(path, paras, font_name, footer_marker):
    d = Document(path)
    idx = next(i for i, p in enumerate(d.paragraphs) if footer_marker in p.text)
    ref = d.paragraphs[idx]
    for text, bold in paras:
        insert_before(ref, text, bold, font_name)
    d.save(path)
    print("saved", path, "| inserted before para", idx, repr(footer_marker[:30]))


add_section("R22-DMF微信公众号文章.docx", ZH_PARAS, "Microsoft YaHei",
            "数据：Agarwal et al., 1982")
add_section("R22-DMF-English-Article.docx", EN_PARAS, None,
            "Data: Agarwal et al., 1982")
